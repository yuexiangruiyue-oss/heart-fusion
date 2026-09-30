# -*- coding: utf-8 -*-
"""批量提取 爱的文章 全部文件内容：docx/pdf/html/xlsx → JSONL"""
import os, json, re, sys
from docx import Document
from pypdf import PdfReader
import openpyxl

SRC = r"D:\双生天使的怀抱\爱的文章"
OUT = r"D:\双生天使的怀抱\.workbuddy\tmp_articles_extract.jsonl"

def ex_docx(p):
    try:
        d = Document(p)
        t = " ".join(x.text.strip() for x in d.paragraphs if x.text.strip())
        return t, None
    except Exception as e:
        return "", f"docx:{str(e)[:100]}"

def ex_pdf(p):
    try:
        r = PdfReader(p)
        t = " ".join((pg.extract_text() or "") for pg in r.pages)
        return " ".join(t.split()), f"pages:{len(r.pages)}"
    except Exception as e:
        return "", f"pdf:{str(e)[:100]}"

def ex_html(p):
    try:
        raw = open(p, "rb").read()
        for enc in ("utf-8", "gbk"):
            try:
                s = raw.decode(enc); break
            except Exception: continue
        s = re.sub(r"<script[\s\S]*?</script>|<style[\s\S]*?</style>", " ", s)
        s = re.sub(r"<[^>]+>", " ", s)
        return " ".join(s.split()), None
    except Exception as e:
        return "", f"html:{str(e)[:100]}"

def ex_xlsx(p):
    try:
        wb = openpyxl.load_workbook(p, read_only=True)
        parts = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None]
                if vals: parts.append(" | ".join(vals))
        return " ".join(parts)[:60000], None
    except Exception as e:
        return "", f"xlsx:{str(e)[:100]}"

files = sorted(os.listdir(SRC))
n = err = 0
with open(OUT, "w", encoding="utf-8") as fo:
    for name in files:
        p = os.path.join(SRC, name)
        if not os.path.isfile(p): continue
        ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
        if ext == "docx": t, info = ex_docx(p)
        elif ext == "pdf": t, info = ex_pdf(p)
        elif ext in ("html", "htm"): t, info = ex_html(p)
        elif ext == "xlsx": t, info = ex_xlsx(p)
        else: t, info = "", f"skip:{ext}"
        note = None
        if not t and info and info.startswith(("docx:","pdf:","html:","xlsx:")):
            note = info; err += 1
        elif info and info.startswith("pages:"):
            note = info
        fo.write(json.dumps({"name": name, "chars": len(t), "preview": t[:500], "info": note}, ensure_ascii=False) + "\n")
        n += 1
        if n % 20 == 0: print(f"{n}/{len(files)}", file=sys.stderr, flush=True)
print(f"DONE files={n} errors={err}")
